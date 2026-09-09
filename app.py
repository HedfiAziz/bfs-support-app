import os
import csv
from io import StringIO, BytesIO
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, session, Response, send_file
from flask_sqlalchemy import SQLAlchemy
from sqlalchemy import text
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

load_dotenv()

app = Flask(__name__)
app.config['SECRET_KEY'] = os.getenv('SECRET_KEY', 'cle_secrete_bfs_2026')
ADMIN_PASSWORD = os.getenv('ADMIN_PASSWORD', 'adminbfs2026')

# Dossier d'enregistrement des pièces jointes
UPLOAD_FOLDER = os.path.join(app.root_path, 'static', 'uploads')
os.makedirs(UPLOAD_FOLDER, exist_ok=True)
app.config['UPLOAD_FOLDER'] = UPLOAD_FOLDER

# Connexion Neon PostgreSQL
db_url = os.getenv('DATABASE_URL')
if db_url and db_url.startswith("postgres://"):
    db_url = db_url.replace("postgres://", "postgresql://", 1)

app.config['SQLALCHEMY_DATABASE_URI'] = db_url
app.config['SQLALCHEMY_TRACK_MODIFICATIONS'] = False

# Options du moteur SQLAlchemy contre la déconnexion SSL inattendue
app.config['SQLALCHEMY_ENGINE_OPTIONS'] = {
    "pool_pre_ping": True,
    "pool_recycle": 300,
}

db = SQLAlchemy(app)

class Ticket(db.Model):
    __tablename__ = 'tickets'

    id = db.Column(db.Integer, primary_key=True)
    entreprise = db.Column(db.String(150), nullable=False)
    nom_prenom = db.Column(db.String(150), nullable=False)
    telephone = db.Column(db.String(50), nullable=False)
    email = db.Column(db.String(150), nullable=False)
    module = db.Column(db.String(255), nullable=False)
    urgence = db.Column(db.String(50), nullable=False)
    sujet = db.Column(db.String(200), nullable=False)
    description = db.Column(db.Text, nullable=False)
    fichier_joint = db.Column(db.String(255), nullable=True)
    temps_estime = db.Column(db.Float, default=0.0)
    statut = db.Column(db.String(50), default="Nouveau")
    date_creation = db.Column(db.DateTime, default=datetime.utcnow)

# Migration automatique PostgreSQL au démarrage
with app.app_context():
    db.create_all()
    try:
        db.session.execute(text("ALTER TABLE tickets ADD COLUMN IF NOT EXISTS fichier_joint VARCHAR(255);"))
        db.session.execute(text("ALTER TABLE tickets ADD COLUMN IF NOT EXISTS temps_estime FLOAT DEFAULT 0.0;"))
        db.session.execute(text("ALTER TABLE tickets ALTER COLUMN module TYPE VARCHAR(255);"))
        db.session.commit()
    except Exception as e:
        db.session.rollback()

@app.route('/')
def index():
    return render_template('index.html')

@app.route('/submit', methods=['POST'])
def submit():
    if request.method == 'POST':
        filename = None
        if 'fichier_joint' in request.files:
            file = request.files['fichier_joint']
            if file and file.filename != '':
                filename = secure_filename(file.filename)
                file.save(os.path.join(app.config['UPLOAD_FOLDER'], filename))

        nouveau_ticket = Ticket(
            entreprise=request.form.get('entreprise', '').strip(),
            nom_prenom=request.form.get('nom_prenom', '').strip(),
            telephone=request.form.get('telephone', '').strip(),
            email=request.form.get('email', '').strip(),
            module=request.form.get('module', '').strip(),
            urgence=request.form.get('urgence', '').strip(),
            sujet=request.form.get('sujet', '').strip(),
            description=request.form.get('description', '').strip(),
            fichier_joint=filename
        )
        db.session.add(nouveau_ticket)
        db.session.commit()

        return render_template('success.html', ticket=nouveau_ticket)

@app.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    error = None
    if request.method == 'POST':
        if request.form.get('password') == ADMIN_PASSWORD:
            session['admin_logged_in'] = True
            return redirect(url_for('admin_dashboard'))
        else:
            error = "Mot de passe incorrect."
    return render_template('admin_login.html', error=error)

@app.route('/admin')
def admin_dashboard():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
    
    tickets = Ticket.query.order_by(Ticket.date_creation.desc()).all()
    
    stats = {
        'total': len(tickets),
        'nouveaux': sum(1 for t in tickets if t.statut == 'Nouveau'),
        'en_cours': sum(1 for t in tickets if t.statut in ['En cours', 'En Cours']),
        'critiques': sum(1 for t in tickets if t.urgence == 'Critique'),
        'resolus': sum(1 for t in tickets if t.statut in ['Résolu', 'Clôturé', 'Cloture', 'Resolue'])
    }

    return render_template('admin.html', tickets=tickets, stats=stats)

@app.route('/admin/update_status/<int:ticket_id>', methods=['POST'])
def update_status(ticket_id):
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
    
    ticket = Ticket.query.get_or_404(ticket_id)
    nouveau_statut = request.form.get('statut')
    if nouveau_statut:
        ticket.statut = nouveau_statut
        db.session.commit()
        
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/update_temps/<int:ticket_id>', methods=['POST'])
def update_temps(ticket_id):
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
    
    ticket = Ticket.query.get_or_404(ticket_id)
    temps_val = request.form.get('temps_estime')
    if temps_val is not None:
        try:
            ticket.temps_estime = float(temps_val)
            db.session.commit()
        except ValueError:
            pass
            
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/delete/<int:ticket_id>', methods=['POST'])
def delete_ticket(ticket_id):
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
    
    ticket = Ticket.query.get_or_404(ticket_id)
    
    if ticket.fichier_joint:
        file_path = os.path.join(app.config['UPLOAD_FOLDER'], ticket.fichier_joint)
        if os.path.exists(file_path):
            try:
                os.remove(file_path)
            except OSError:
                pass

    db.session.delete(ticket)
    db.session.commit()
    return redirect(url_for('admin_dashboard'))

@app.route('/admin/export_csv')
def export_csv():
    if not session.get('admin_logged_in'):
        return redirect(url_for('admin_login'))
    
    tickets = Ticket.query.order_by(Ticket.date_creation.desc()).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Suivi des Tickets BFS"
    ws.views.sheetView[0].showGridLines = True

    # Couleurs de charte graphique BFS
    navy_fill = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
    card_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
    header_fill = PatternFill(start_color="2C4D75", end_color="2C4D75", fill_type="solid")
    
    # Styles de police
    title_font = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    card_title_font = Font(name="Calibri", size=9, bold=True, color="555555")
    card_val_font = Font(name="Calibri", size=14, bold=True, color="1B365D")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)

    thin_border = Border(
        left=Side(style='thin', color='D3D3D3'),
        right=Side(style='thin', color='D3D3D3'),
        top=Side(style='thin', color='D3D3D3'),
        bottom=Side(style='thin', color='D3D3D3')
    )

    # 1. Bandeau de Titre
    ws.merge_cells("A1:M1")
    title_cell = ws["A1"]
    title_cell.value = "  BFS SOLUTIONS — TABLEAU DE BORD & SUIVI DES INCIDENTS"
    title_cell.font = title_font
    title_cell.fill = navy_fill
    title_cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 40

    # 2. Calculs pour le Mini Dashboard
    total_tickets = len(tickets)
    nouveaux = sum(1 for t in tickets if t.statut == 'Nouveau')
    en_cours = sum(1 for t in tickets if t.statut in ['En cours', 'En Cours'])
    resolus = sum(1 for t in tickets if t.statut in ['Résolu', 'Clôturé', 'Cloture', 'Resolue'])
    temps_cumule = sum(t.temps_estime or 0.0 for t in tickets)

    kpis = [
        ("TOTAL TICKETS", total_tickets, "B3", "B4"),
        ("NOUVEAUX", nouveaux, "D3", "D4"),
        ("EN COURS", en_cours, "F3", "F4"),
        ("RÉSOLUS", resolus, "H3", "H4"),
        ("HEURES PASSÉES", f"{temps_cumule:.1f} h", "J3", "J4")
    ]

    for label, value, c_lbl, c_val in kpis:
        cell_lbl = ws[c_lbl]
        cell_lbl.value = label
        cell_lbl.font = card_title_font
        cell_lbl.fill = card_fill
        cell_lbl.alignment = Alignment(horizontal="center", vertical="center")

        cell_val = ws[c_val]
        cell_val.value = value
        cell_val.font = card_val_font
        cell_val.fill = card_fill
        cell_val.alignment = Alignment(horizontal="center", vertical="center")

    # 3. En-têtes du Tableau (Ligne 6)
    headers = [
        'ID', 'Date Création', 'Entreprise', 'Demandeur', 'Téléphone', 
        'Email', 'Module', 'Urgence', 'Temps Passé (h)', 'Sujet', 
        'Description', 'Fichier Joint', 'Statut'
    ]
    
    header_row = 6
    ws.row_dimensions[header_row].height = 25
    
    for col_idx, text_header in enumerate(headers, 1):
        c = ws.cell(row=header_row, column=col_idx)
        c.value = text_header
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = thin_border

    # 4. Données des Tickets
    start_data_row = 7
    for t in tickets:
        row_vals = [
            t.id,
            t.date_creation.strftime('%d/%m/%Y %H:%M') if t.date_creation else '',
            t.entreprise,
            t.nom_prenom,
            t.telephone,
            t.email,
            t.module,
            t.urgence,
            t.temps_estime or 0.0,
            t.sujet,
            t.description,
            t.fichier_joint or '',
            t.statut
        ]
        
        ws.row_dimensions[start_data_row].height = 20
        for col_idx, val in enumerate(row_vals, 1):
            c = ws.cell(row=start_data_row, column=col_idx)
            c.value = val
            c.font = data_font
            c.border = thin_border
            
            # Alignements ciblés
            if col_idx in [1, 2, 8, 9, 13]:
                c.alignment = Alignment(horizontal="center", vertical="center")
            else:
                c.alignment = Alignment(horizontal="left", vertical="center")

        start_data_row += 1

    # 5. Ajustement automatique de la largeur des colonnes
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row < 6:
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 50)

    output = BytesIO()
    wb.save(output)
    output.seek(0)

    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"Tickets_BFS_Solutions_{datetime.now().strftime('%Y%m%d')}.xlsx"
    )

@app.route('/admin/logout')
def admin_logout():
    session.pop('admin_logged_in', None)
    return redirect(url_for('admin_login'))

if __name__ == '__main__':
    app.run(debug=True)